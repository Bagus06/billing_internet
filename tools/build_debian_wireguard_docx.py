from __future__ import print_function

import os
import re
import sys
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
SOURCE = os.path.join(ROOT, 'docs', 'DEBIAN_WIREGUARD_RUNBOOK.md')
OUT_DIR = os.path.join(ROOT, 'build', 'debian-wireguard-docx')
DOCX_PATH = os.path.join(OUT_DIR, 'Panduan_Konfigurasi_Debian_WireGuard_BATARA.docx')
FLOW_PATH = os.path.join(OUT_DIR, 'arsitektur-koneksi.png')

BLUE = '155D96'
BLUE_DARK = '0C3657'
BLUE_LIGHT = 'EAF3FA'
CYAN = '20A4C7'
GRAY = '52606D'
GRAY_LIGHT = 'F4F7FA'
WHITE = 'FFFFFF'
RED = 'B42318'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + margin))
        if node is None:
            node = OxmlElement('w:' + margin)
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_begin, instr, fld_end])


def add_bottom_border(paragraph, color=CYAN, size='18'):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        p_bdr = OxmlElement('w:pBdr')
        p_pr.append(p_bdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '5')
    bottom.set(qn('w:color'), color)
    p_bdr.append(bottom)


def add_inline_runs(paragraph, text, color=None):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
            run.font.highlight_color = None
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        else:
            run = paragraph.add_run(part)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def font(size, bold=False):
    paths = [
        'C:/Windows/Fonts/arial.ttf',
        'C:/Windows/Fonts/calibri.ttf',
    ]
    if bold:
        paths = ['C:/Windows/Fonts/arialbd.ttf', 'C:/Windows/Fonts/calibrib.ttf'] + paths
    for path in paths:
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, subtitle, fill, outline, icon=None):
    x1, y1, x2, y2 = xy
    draw.rectangle(xy, fill=fill, outline=outline, width=3)
    title_font = font(28, True)
    sub_font = font(20)
    draw.text((x1 + 28, y1 + 26), title, font=title_font, fill=(14, 54, 87))
    lines = subtitle.split('\n')
    yy = y1 + 72
    for line in lines:
        draw.text((x1 + 28, yy), line, font=sub_font, fill=(63, 78, 92))
        yy += 28


def arrow(draw, start, end, label):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=(32, 164, 199), width=5)
    angle = 13
    draw.polygon([(x2, y2), (x2 - 18, y2 - angle), (x2 - 18, y2 + angle)], fill=(32, 164, 199))
    if label:
        f = font(17, True)
        w, _ = draw.textsize(label, font=f)
        cx = (x1 + x2) // 2
        cy = (y1 + y2) // 2
        draw.rectangle((cx - w // 2 - 9, cy - 17, cx + w // 2 + 9, cy + 12), fill=(255, 255, 255))
        draw.text((cx - w // 2, cy - 13), label, font=f, fill=(12, 54, 87))


def create_flowchart(path):
    img = Image.new('RGB', (1800, 980), (247, 250, 252))
    draw = ImageDraw.Draw(img)
    draw.text((70, 38), 'Arsitektur Koneksi BATARA', font=font(40, True), fill=(12, 54, 87))
    draw.text((70, 94), 'Akses aman, layanan isolir, dan relay data OLT', font=font(23), fill=(82, 96, 109))

    rounded_box(draw, (70, 190, 410, 350), 'Laptop Admin', 'WireGuard client\n10.88.0.2/32', (234, 243, 250), (21, 93, 150))
    rounded_box(draw, (590, 190, 930, 350), 'MikroTik', 'Public gateway\nNAT & firewall', (226, 246, 249), (32, 164, 199))
    rounded_box(draw, (1110, 150, 1630, 390), 'Server Debian', 'WireGuard 10.88.0.1\nSSH :30003 | Nginx :80\nOLT Relay :8787', (234, 243, 250), (21, 93, 150))
    arrow(draw, (410, 270), (590, 270), 'UDP 51888')
    arrow(draw, (930, 270), (1110, 270), 'DST-NAT')

    rounded_box(draw, (70, 560, 410, 720), 'Pelanggan Isolir', 'Pool 10.7.0.0/29\nHTTP redirect', (255, 244, 229), (217, 119, 6))
    arrow(draw, (410, 640), (590, 640), 'HTTP')
    rounded_box(draw, (590, 560, 930, 720), 'NAT / Firewall', 'Arahkan trafik web\nke Debian', (226, 246, 249), (32, 164, 199))
    arrow(draw, (930, 640), (1110, 640), 'TCP 80')
    rounded_box(draw, (1110, 520, 1630, 760), 'Layanan Internal', 'Halaman isolir (Nginx)\nSSH administrasi\nRelay SNMP dan cache', (234, 243, 250), (21, 93, 150))

    rounded_box(draw, (1110, 820, 1450, 940), 'OLT HSGQ', '192.168.99.1\nSNMP UDP 161', (239, 247, 237), (39, 135, 74))
    draw.line((1370, 760, 1370, 820), fill=(39, 135, 74), width=5)
    draw.polygon([(1370, 820), (1357, 800), (1383, 800)], fill=(39, 135, 74))
    img.save(path, quality=95)


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string('243746')
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.16

    for name, size, color, before, after in [
        ('Title', 30, BLUE_DARK, 0, 12),
        ('Heading 1', 17, BLUE, 18, 8),
        ('Heading 2', 13.5, BLUE_DARK, 14, 6),
        ('Heading 3', 11.5, BLUE_DARK, 10, 4),
    ]:
        s = styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    for name in ['List Bullet', 'List Number']:
        s = styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(10.5)
        s.paragraph_format.left_indent = Inches(0.28)
        s.paragraph_format.first_line_indent = Inches(-0.18)
        s.paragraph_format.space_after = Pt(3)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run('BATARA  |  INFRASTRUCTURE RUNBOOK')
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    add_bottom_border(header, color='B9D9ED', size='8')

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('Dokumen operasional internal  •  Halaman ')
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(footer)
    return doc


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('BATARA')
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(CYAN)

    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    p.add_run('Panduan Konfigurasi\nDebian & WireGuard')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Server isolir • administrasi aman • relay OLT')
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('RUNBOOK IMPLEMENTASI & MIGRASI')
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), BLUE)
    p_pr.append(shd)

    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Inches(1.55)
    info.columns[1].width = Inches(4.3)
    rows = [
        ('Dokumen', 'Panduan instalasi ulang dan migrasi server BATARA'),
        ('Platform', 'Debian 13 minimal, MikroTik, WireGuard, Nginx'),
        ('Versi', '1.0'),
        ('Diperbarui', datetime.now().strftime('%d %B %Y')),
    ]
    for idx, (label, value) in enumerate(rows):
        for cell in info.rows[idx].cells:
            set_cell_margins(cell, top=100, bottom=100)
        set_cell_shading(info.cell(idx, 0), BLUE_LIGHT)
        info.cell(idx, 0).paragraphs[0].add_run(label).bold = True
        info.cell(idx, 1).paragraphs[0].add_run(value)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('PERINGATAN KEAMANAN')
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(RED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Gunakan placeholder untuk password, private key, token API, dan community SNMP. Jangan simpan rahasia aktual di dokumen atau Git.')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_page_break()


def add_contents(doc, headings):
    p = doc.add_paragraph('Daftar Isi', style='Heading 1')
    add_bottom_border(p)
    for level, title in headings:
        if level != 2:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    doc.add_page_break()


def add_code_block(doc, lines, language):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.55)
    cell = table.cell(0, 0)
    set_cell_shading(cell, '132A3A')
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if language:
        r = p.add_run(language.upper() + '\n')
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string('6ED3EA')
    r = p.add_run('\n'.join(lines).rstrip())
    r.font.name = 'Consolas'
    r.font.size = Pt(7.8)
    r.font.color.rgb = RGBColor.from_string('F3F7FA')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown_table(doc, rows):
    if len(rows) < 2:
        return
    parsed = [[c.strip() for c in row.strip().strip('|').split('|')] for row in rows]
    if len(parsed) > 1 and all(re.match(r'^:?-{3,}:?$', c) for c in parsed[1]):
        parsed.pop(1)
    cols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = 6.55
    widths = [total / cols] * cols
    if cols == 2:
        widths = [2.05, 4.5]
    for ri, row in enumerate(parsed):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, BLUE)
            elif ri % 2 == 0:
                set_cell_shading(cell, GRAY_LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text = row[ci] if ci < len(row) else ''
            add_inline_runs(p, text, WHITE if ri == 0 else None)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc, lines):
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith('# '):
            i += 1
            continue
        if line.startswith('```'):
            language = line[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i].rstrip('\n'))
                i += 1
            i += 1
            if language == 'mermaid':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(FLOW_PATH, width=Inches(6.55))
                cap = doc.add_paragraph('Gambar 1. Alur koneksi administrasi, isolir, dan relay OLT')
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(8)
                cap.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
            else:
                add_code_block(doc, block, language)
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, rows)
            continue
        heading = re.match(r'^(#{2,3})\s+(.+)$', line)
        if heading:
            level = len(heading.group(1)) - 1
            p = doc.add_paragraph(heading.group(2), style='Heading {}'.format(level))
            if level == 1:
                add_bottom_border(p)
            i += 1
            continue
        bullet = re.match(r'^[-*]\s+(.+)$', line)
        number = re.match(r'^\d+\.\s+(.+)$', line)
        if bullet or number:
            p = doc.add_paragraph(style='List Bullet' if bullet else 'List Number')
            add_inline_runs(p, (bullet or number).group(1))
            i += 1
            continue
        quote = re.match(r'^>\s*(.+)$', line)
        if quote:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.24)
            p.paragraph_format.right_indent = Inches(0.1)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), BLUE_LIGHT)
            p_pr.append(shd)
            add_inline_runs(p, quote.group(1))
            i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or nxt.startswith(('#', '```', '|', '-', '*', '>')) or re.match(r'^\d+\.\s+', nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline_runs(p, ' '.join(paragraph_lines))


def main():
    if not os.path.exists(OUT_DIR):
        os.makedirs(OUT_DIR)
    with open(SOURCE, 'r', encoding='utf-8') as fh:
        content = fh.read()
    lines = content.splitlines()
    headings = []
    for line in lines:
        m = re.match(r'^(#{2,3})\s+(.+)$', line)
        if m:
            headings.append((len(m.group(1)), m.group(2)))

    create_flowchart(FLOW_PATH)
    doc = setup_document()
    add_cover(doc)
    add_contents(doc, headings)
    add_body(doc, lines)

    props = doc.core_properties
    props.title = 'Panduan Konfigurasi Debian & WireGuard BATARA'
    props.subject = 'Runbook server isolir, administrasi WireGuard, dan relay OLT'
    props.author = 'BATARA'
    props.keywords = 'Debian, WireGuard, MikroTik, isolir, OLT, SNMP, Nginx'
    props.comments = 'Dokumen dibuat dari docs/DEBIAN_WIREGUARD_RUNBOOK.md'
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == '__main__':
    main()
